"""Generates the Phase III presentation deck for EcoMesh, matching the
visual template of the Phase II deck (dark background, coral title accent,
cyan highlights, numbered section eyebrows, card/table layouts)."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

OUT = "Phase_III_EcoMesh_Group20.pptx"

# --- palette, matching the Phase II deck -----------------------------------------
BG = RGBColor(0x0B, 0x10, 0x1A)
CARD_BG = RGBColor(0x14, 0x1C, 0x2B)
CARD_BORDER = RGBColor(0x2A, 0x35, 0x48)
CORAL = RGBColor(0xFF, 0x5A, 0x5F)
CYAN = RGBColor(0x5E, 0xEA, 0xD4)
WHITE = RGBColor(0xF2, 0xF4, 0xF8)
GREY = RGBColor(0x9A, 0xA4, 0xB5)
TABLE_HEAD = RGBColor(0x5E, 0xEA, 0xD4)
TABLE_HEAD_TEXT = RGBColor(0x0B, 0x10, 0x1A)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

TOTAL_SLIDES = 12


def new_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def blank_slide(prs):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    # send to back
    spTree = slide.shapes._spTree
    spTree.remove(bg._element)
    spTree.insert(2, bg._element)
    return slide


def add_text(slide, left, top, width, height, text, size, color, bold=False,
             italic=False, align=PP_ALIGN.LEFT, font="Arial", anchor=None, line_spacing=1.0):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    if anchor is not None:
        tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    return box


def add_multiline(slide, left, top, width, height, lines, size, color, bold=False,
                   align=PP_ALIGN.LEFT, font="Arial", space_after=4, line_spacing=1.1):
    """lines: list of (text, bold_override) or plain strings."""
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(lines):
        text, b = item if isinstance(item, tuple) else (item, bold)
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = b
        run.font.name = font
    return box


def card(slide, left, top, width, height, fill=CARD_BG, border=CARD_BORDER):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = border
    box.line.width = Pt(0.75)
    box.shadow.inherit = False
    return box


def chrome(slide, eyebrow, title, page_num, subtitle=None):
    """Standard header (eyebrow + title [+ subtitle]) and footer, matching
    the Phase II deck's layout exactly."""
    add_text(slide, Inches(0.55), Inches(0.35), Inches(8), Inches(0.35),
              eyebrow, 11, GREY, bold=True)
    add_text(slide, Inches(0.55), Inches(0.72), Inches(11.5), Inches(0.7),
              title, 28, WHITE, bold=False)
    y = Inches(1.4)
    if subtitle:
        add_text(slide, Inches(0.55), Inches(1.3), Inches(11.5), Inches(0.55),
                  subtitle, 14, CYAN)
        y = Inches(1.95)
    # footer divider
    line = slide.shapes.add_connector(1, Inches(0.55), Inches(7.05), Inches(12.78), Inches(7.05))
    line.line.color.rgb = CARD_BORDER
    line.line.width = Pt(0.75)
    add_text(slide, Inches(0.55), Inches(7.12), Inches(4), Inches(0.3),
              "ECOMESH   /   GROUP 20", 9, GREY, bold=True)
    add_text(slide, Inches(10.5), Inches(7.12), Inches(2.28), Inches(0.3),
              f"{page_num:02d} / {TOTAL_SLIDES}", 9, GREY, bold=True, align=PP_ALIGN.RIGHT)
    return y


def eyebrow_label(n, label):
    return f"{n:02d} · {label}"


def build():
    prs = new_presentation()

    # ================================================================ Slide 1: Title
    s = blank_slide(prs)
    add_text(s, Inches(0.7), Inches(0.6), Inches(8), Inches(0.4),
              "MINI PROJECT II · GROUP 20", 13, WHITE, bold=True)
    add_text(s, Inches(0.68), Inches(1.5), Inches(10), Inches(1.3),
              "EcoMesh", 60, CORAL, bold=True)
    add_text(s, Inches(0.7), Inches(2.75), Inches(10.5), Inches(0.9),
              "Distributed Peer-to-Peer Disaster\nRelief Communication Mesh", 22, WHITE, line_spacing=1.15)
    add_text(s, Inches(0.7), Inches(3.95), Inches(10.5), Inches(0.4),
              "PHASE III — ROUTING, RELIABILITY & SECURITY", 15, CORAL, bold=True)
    line = s.shapes.add_connector(1, Inches(0.7), Inches(4.5), Inches(12.6), Inches(4.5))
    line.line.color.rgb = CARD_BORDER
    line.line.width = Pt(1)
    add_text(s, Inches(0.7), Inches(4.75), Inches(3), Inches(0.3), "TEAM", 11, GREY, bold=True)
    add_text(s, Inches(0.7), Inches(5.05), Inches(10), Inches(0.4),
              "Miti Shah  ·  Twissha Shah  ·  Tanvi Sheth", 16, WHITE)
    add_text(s, Inches(0.7), Inches(5.6), Inches(3), Inches(0.3), "GUIDE", 11, GREY, bold=True)
    add_text(s, Inches(1.65), Inches(5.57), Inches(6), Inches(0.4), "Prof. Anuj Tawari", 16, WHITE)

    # ================================================================ Slide 2: Recap
    s = blank_slide(prs)
    chrome(s, eyebrow_label(2, "RECAP"), "From discovery to a routable, trustworthy mesh", 2)
    add_text(s, Inches(0.55), Inches(1.35), Inches(11.8), Inches(0.6),
              "Phase II proved two nodes could discover each other and relay a single hop. Phase III turns "
              "that into a real mesh — many hops, the best path, nothing lost, nothing readable by a stranger.",
              15, CYAN, line_spacing=1.2)
    left_card = card(s, Inches(0.55), Inches(2.3), Inches(5.95), Inches(4.3))
    add_text(s, Inches(0.85), Inches(2.55), Inches(5.4), Inches(0.35), "WHAT PHASE II LEFT OFF", 12, CYAN, bold=True)
    add_multiline(s, Inches(0.85), Inches(3.0), Inches(5.4), Inches(3.4), [
        "Nodes could discover direct neighbours only.",
        "A message could be relayed, but only one hop.",
        "No weighting — every link treated as equally good.",
        "No encryption, no signing, no persistent log.",
    ], 14, WHITE, space_after=14)
    right_card = card(s, Inches(6.8), Inches(2.3), Inches(5.95), Inches(4.3))
    add_text(s, Inches(7.1), Inches(2.55), Inches(5.4), Inches(0.35), "WHAT PHASE III ADDS", 12, CORAL, bold=True)
    add_multiline(s, Inches(7.1), Inches(3.0), Inches(5.4), Inches(3.4), [
        "Network-wide routing — every node learns the full graph.",
        "Weighted paths — reliability beats hop count.",
        "Reliable delivery — priority, retry, store-and-forward.",
        "Security — end-to-end encryption and signing.",
        "Persistent logging — every message and event, queryable.",
    ], 14, WHITE, space_after=10)

    # ================================================================ Slide 3: Phase 3 at a glance
    s = blank_slide(prs)
    chrome(s, eyebrow_label(3, "SCOPE"), "Phase III at a glance", 3)
    items = [
        ("01", "ADAPTIVE ROUTING", ["Link-state gossip,", "re-flooded network-wide", "Weighted Dijkstra over", "measured link quality"]),
        ("02", "RELIABILITY", ["Priority queue, duplicate", "suppression", "ACK + retry, bounded", "store-and-forward"]),
        ("03", "SECURITY", ["AES-256-GCM payload", "encryption", "Ed25519 header signing,", "key distribution"]),
        ("04", "STORAGE", ["SQLite (WAL) message", "and event log", "Every send, forward,", "drop, retry — queryable"]),
    ]
    cw = Inches(2.9)
    gap = Inches(0.2)
    x0 = Inches(0.55)
    for i, (num, label, lines) in enumerate(items):
        x = Emu(x0 + i * (cw + gap))
        c = card(s, x, Inches(2.0), cw, Inches(4.5))
        add_text(s, Emu(x + Inches(0.25)), Inches(2.25), Inches(2.4), Inches(0.6), num, 26, CORAL, bold=True)
        add_text(s, Emu(x + Inches(0.25)), Inches(2.95), Inches(2.4), Inches(0.5), label, 12.5, CYAN, bold=True)
        add_multiline(s, Emu(x + Inches(0.25)), Inches(3.45), Inches(2.4), Inches(2.8), lines, 13, WHITE, space_after=8)

    # ================================================================ Slide 4: Four fixes that matter
    s = blank_slide(prs)
    chrome(s, eyebrow_label(4, "ENGINEERING"), "Four bugs that would have failed silently", 4,
            subtitle="Found by writing the tests, not by guessing — each one is now a regression test.")
    fixes = [
        ("GRAPH STUCK AT TWO HOPS", "Link-state reports weren't re-flooded past the first neighbour — a node "
         "could never learn a route beyond 2 hops. Fixed with origin + sequence numbers so reports "
         "propagate network-wide and stale replays are rejected."),
        ("FAILURES NEVER PENALISED", "The retry logic never told routing.py about a timed-out link — the cost "
         "function's failure term was always zero. Now every ACK timeout updates the link's cost "
         "immediately, not on the next gossip cycle."),
        ("RETENTION QUEUE UNBOUNDED", "Store-and-forward had no size or age limit. Now it evicts the lowest-"
         "priority message first (SOS is never displaced) and expires anything held past 30 minutes."),
        ("SIGNATURES WOULD BREAK AT HOP 1", "The spec says “sign header + path” — but ttl, hop_count and path "
         "all change at every relay. Signing them would make every signature invalid after one hop. "
         "Fixed by signing only the fields that never change in transit."),
    ]
    cw = Inches(5.95)
    ch = Inches(2.15)
    gap = Inches(0.2)
    x0 = Inches(0.55)
    y0 = Inches(2.1)
    for i, (title, body) in enumerate(fixes):
        col = i % 2
        row = i // 2
        x = Emu(x0 + col * (cw + gap))
        y = Emu(y0 + row * (ch + gap))
        card(s, x, y, cw, ch)
        add_text(s, Emu(x + Inches(0.3)), Emu(y + Inches(0.2)), Emu(cw - Inches(0.6)), Inches(0.4),
                  title, 13, CORAL, bold=True)
        add_text(s, Emu(x + Inches(0.3)), Emu(y + Inches(0.65)), Emu(cw - Inches(0.6)), Emu(ch - Inches(0.85)),
                  body, 11.5, WHITE, line_spacing=1.2)

    # ================================================================ Slide 5: Weighted routing
    s = blank_slide(prs)
    chrome(s, eyebrow_label(5, "ALGORITHM · ROUTING"), "A longer path can be the better path", 5)
    add_text(s, Inches(0.55), Inches(1.5), Inches(11.5), Inches(0.35), "WEIGHTED LINK COST", 12, CORAL, bold=True)
    formula_card = card(s, Inches(0.55), Inches(1.9), Inches(12.2), Inches(0.95))
    add_text(s, Inches(0.55), Inches(2.15), Inches(12.2), Inches(0.5),
              "cost(u,v) = H + α × (1 − quality) + β × failures", 20, CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.55), Inches(3.15), Inches(11.5), Inches(0.35), "WORKED COMPARISON", 12, CORAL, bold=True)

    tbl_left, tbl_top, tbl_w, tbl_h = Inches(0.55), Inches(3.55), Inches(8.5), Inches(1.5)
    gt = s.shapes.add_table(3, 4, tbl_left, tbl_top, tbl_w, tbl_h).table
    headers = ["PATH", "COST", "DECISION", "STATUS"]
    rows = [
        ("2 hops · poor links", "7.75", "REJECT", "—"),
        ("3 hops · good links", "3.75", "CHOOSE", "✓ tested"),
    ]
    for j, h in enumerate(headers):
        cell = gt.cell(0, j)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = TABLE_HEAD
        run = cell.text_frame.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = TABLE_HEAD_TEXT
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = gt.cell(i + 1, j)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            run = cell.text_frame.paragraphs[0].runs[0]
            run.font.size = Pt(12)
            run.font.color.rgb = CORAL if val == "REJECT" else (CYAN if val in ("CHOOSE", "✓ tested") else WHITE)

    add_text(s, Inches(0.55), Inches(5.35), Inches(11.8), Inches(1.4),
              "Dijkstra over the weighted graph · O((V+E) log V). Proven directly by an automated test: "
              "a two-hop path through unreliable links is rejected in favour of a three-hop path through "
              "good ones — the headline claim of the whole routing design, not just a formula on a slide.",
              14, WHITE, line_spacing=1.25)

    # ================================================================ Slide 6: Priority
    s = blank_slide(prs)
    chrome(s, eyebrow_label(6, "ALGORITHM · PRIORITY"), "SOS moves to the front — at every hop", 6,
            subtitle="Binary min-heap · ordered by (priority, creation time)")
    labels = [("NORMAL", "1"), ("STATUS", "2"), ("NORMAL", "1")]
    x = Inches(0.55)
    for name, val in labels:
        c = card(s, x, Inches(2.6), Inches(2.1), Inches(1.0))
        add_text(s, Emu(x + Inches(0.2)), Inches(2.9), Inches(1.0), Inches(0.4), name, 14, WHITE, bold=True)
        add_text(s, Emu(x + Inches(1.5)), Inches(2.9), Inches(0.5), Inches(0.4), val, 14, GREY, bold=True)
        x = Emu(x + Inches(2.3))
    add_text(s, Inches(7.15), Inches(2.95), Inches(0.6), Inches(0.5), "→", 24, GREY, bold=True)
    sos_card = card(s, Inches(7.85), Inches(2.6), Inches(2.4), Inches(1.0), fill=CORAL, border=CORAL)
    add_text(s, Inches(8.05), Inches(2.9), Inches(1.2), Inches(0.4), "SOS", 15, WHITE, bold=True)
    add_text(s, Inches(9.6), Inches(2.9), Inches(0.5), Inches(0.4), "0", 15, WHITE, bold=True)
    add_text(s, Inches(0.55), Inches(4.0), Inches(11.8), Inches(0.5),
              "SELECTED NEXT: SOS — even though it arrived last", 22, WHITE, bold=True)
    add_multiline(s, Inches(0.55), Inches(4.8), Inches(11.8), Inches(1.8), [
        "Priority is re-applied at every hop, not just at the source — an SOS queued behind NORMAL "
        "traffic at a busy relay still jumps to the front there too.",
        "O(log n) enqueue/dequeue. Proven with a test: a NORMAL packet queued first, an SOS queued "
        "second — the SOS is still the one sent first.",
    ], 14, WHITE, space_after=12, line_spacing=1.2)

    # ================================================================ Slide 7: Security
    s = blank_slide(prs)
    chrome(s, eyebrow_label(7, "ALGORITHM · SECURITY"), "Strangers relay it. They still can't read it.", 7)
    left = card(s, Inches(0.55), Inches(1.8), Inches(5.95), Inches(4.8))
    add_text(s, Inches(0.85), Inches(2.05), Inches(5.4), Inches(0.4), "CONFIDENTIALITY", 13, CORAL, bold=True)
    add_multiline(s, Inches(0.85), Inches(2.55), Inches(5.4), Inches(3.9), [
        "AES-256-GCM payload encryption, end to end.",
        "A relay forwards the ciphertext untouched — it never",
        "calls decrypt, by construction.",
        "",
        "Proven directly: a relay node configured with no key",
        "at all still forwards correctly — and its own log never",
        "records a plaintext body for a message it only relayed.",
    ], 14, WHITE, space_after=6, line_spacing=1.25)
    right = card(s, Inches(6.8), Inches(1.8), Inches(5.95), Inches(4.8))
    add_text(s, Inches(7.1), Inches(2.05), Inches(5.4), Inches(0.4), "AUTHENTICITY", 13, CYAN, bold=True)
    add_multiline(s, Inches(7.1), Inches(2.55), Inches(5.4), Inches(3.9), [
        "Ed25519 header signing. Public keys travel in HELLO,",
        "then propagate network-wide through the same",
        "re-flooding mechanism as routing — a destination three",
        "hops from the sender can still verify it.",
        "",
        "The fix: sign only the fields that survive relaying",
        "(not ttl/hop_count/path). Proven with a 3-hop relay",
        "test — the signature verifies correctly at the far end.",
    ], 14, WHITE, space_after=6, line_spacing=1.25)

    # ================================================================ Slide 8: Code — routing + reliability
    s = blank_slide(prs)
    chrome(s, eyebrow_label(8, "IMPLEMENTATION · PHASE III"), "Routing and reliability are built", 8)
    def module_row(y, name, color, lines):
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), y, Inches(0.07), Inches(1.0))
        bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background(); bar.shadow.inherit = False
        card(s, Inches(0.7), y, Inches(11.65), Inches(1.0))
        add_text(s, Inches(0.95), Emu(y + Inches(0.12)), Inches(2.6), Inches(0.75), name, 16, color, bold=True,
                  anchor=MSO_ANCHOR.MIDDLE)
        add_multiline(s, Inches(3.7), Emu(y + Inches(0.12)), Inches(8.4), Inches(0.75), lines, 12.5, WHITE, space_after=3)

    module_row(Inches(1.9), "routing.py", CYAN, [
        "Link-state gossip with re-flooding · weighted Dijkstra · route caching",
        "Gateway/rescue terminal selection folded directly into route computation",
    ])
    module_row(Inches(3.15), "reliability.py", CYAN, [
        "Priority queue (binary min-heap) · duplicate suppression · ACK tracking with bounded retry",
        "Store-and-forward with count + age bounds and priority-protected eviction",
    ])
    module_row(Inches(4.4), "node.py (wiring)", CORAL, [
        "Node.send() — real message origination with ACK/retry/store-forward lifecycle",
        "Multi-hop forwarding via route_resolver, role propagation through gossip",
    ])
    add_text(s, Inches(0.55), Inches(5.9), Inches(11.8), Inches(0.9),
              "Every routing decision — which link, which hop, which path — now comes from measured "
              "reliability, not assumption.", 14, CYAN, italic=True, line_spacing=1.2)

    # ================================================================ Slide 9: Code — crypto + store
    s = blank_slide(prs)
    chrome(s, eyebrow_label(9, "IMPLEMENTATION · PHASE III"), "Security and storage, wired end to end", 9)
    module_row2_y = Inches(1.9)
    def module_row2(y, name, color, lines):
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), y, Inches(0.07), Inches(1.0))
        bar.fill.solid(); bar.fill.fore_color.rgb = color; bar.line.fill.background(); bar.shadow.inherit = False
        card(s, Inches(0.7), y, Inches(11.65), Inches(1.0))
        add_text(s, Inches(0.95), Emu(y + Inches(0.12)), Inches(2.6), Inches(0.75), name, 16, color, bold=True,
                  anchor=MSO_ANCHOR.MIDDLE)
        add_multiline(s, Inches(3.7), Emu(y + Inches(0.12)), Inches(8.4), Inches(0.75), lines, 12.5, WHITE, space_after=3)

    module_row2(Inches(1.9), "crypto.py", CYAN, [
        "AES-256-GCM encrypt/decrypt · Ed25519 generate/sign/verify",
        "Vetted library primitives only — nothing implemented from first principles",
    ])
    module_row2(Inches(3.15), "store.py", CYAN, [
        "SQLite, write-ahead logging · one database file per node",
        "messages, events, peers — every send/recv/forward/drop/retry, queryable",
    ])

    stat = card(s, Inches(0.55), Inches(4.5), Inches(11.8), Inches(1.3), fill=CYAN, border=CYAN)
    add_text(s, Inches(0.9), Inches(4.75), Inches(2.2), Inches(0.8), "256", 44, BG, bold=True)
    add_text(s, Inches(2.9), Inches(4.8), Inches(5), Inches(0.4), "automated tests passing", 16, BG, bold=True)
    add_text(s, Inches(2.9), Inches(5.25), Inches(8.2), Inches(0.5),
              "up from 79 at the end of Phase II — routing · reliability · crypto · store · node integration",
              12, BG)

    add_text(s, Inches(0.55), Inches(6.1), Inches(11.8), Inches(0.7),
              "Every new feature in Phase III shipped with the test that proves it — including the four "
              "bugs on slide 4, each caught by a test before it ever reached a demo.", 13, GREY, italic=True,
              line_spacing=1.2)

    # ================================================================ Slide 10: Git evidence
    s = blank_slide(prs)
    chrome(s, eyebrow_label(10, "GIT EVIDENCE"), "The repository shows the work", 10,
            subtitle="git log --graph  ·  git shortlog -sn")
    left = card(s, Inches(0.55), Inches(2.1), Inches(5.95), Inches(4.5))
    add_text(s, Inches(0.85), Inches(2.35), Inches(5.4), Inches(0.4), "BRANCH LINEAGE", 12, CORAL, bold=True)
    add_multiline(s, Inches(0.85), Inches(2.85), Inches(5.4), Inches(3.6), [
        "develop → develop_2 → develop_3",
        "Each phase branches from the last, inheriting",
        "everything — nothing re-typed, nothing lost.",
        "",
        "Tagged: v0.1-phase1 · v0.2-phase2 · v0.3-phase3",
        "",
        "Conventional Commits throughout — feat(routing):,",
        "fix(reliability):, docs: — readable without opening",
        "a single file.",
    ], 13.5, WHITE, space_after=8, line_spacing=1.2)
    right = card(s, Inches(6.8), Inches(2.1), Inches(5.95), Inches(4.5))
    add_text(s, Inches(7.1), Inches(2.35), Inches(5.4), Inches(0.4), "WHAT TO RUN LIVE", 12, CYAN, bold=True)
    add_multiline(s, Inches(7.1), Inches(2.85), Inches(5.4), Inches(3.6), [
        "git log --graph --oneline",
        "  — the commit history, phase by phase",
        "",
        "git shortlog -sn",
        "  — commit distribution across the team",
        "",
        "python -m pytest -q",
        "  — 256 passed, live, not a screenshot",
    ], 13.5, WHITE, space_after=10, line_spacing=1.2, font="Consolas")

    # ================================================================ Slide 11: Summary & Next
    s = blank_slide(prs)
    chrome(s, eyebrow_label(11, "SUMMARY & NEXT"), "From a routable mesh to a visible one", 11)
    left = card(s, Inches(0.55), Inches(2.1), Inches(5.95), Inches(4.5))
    add_text(s, Inches(0.85), Inches(2.35), Inches(5.4), Inches(0.4), "DONE · PHASES I–III", 12, CYAN, bold=True)
    add_multiline(s, Inches(0.85), Inches(2.85), Inches(5.4), Inches(3.6), [
        "SRS, six UML diagrams, Design of Algorithms",
        "Protocol, transport, discovery, relay, simulator",
        "Weighted routing, reliability, security, storage —",
        "all wired into Node, 256 tests passing",
    ], 14, WHITE, space_after=14)
    right = card(s, Inches(6.8), Inches(2.1), Inches(5.95), Inches(4.5))
    add_text(s, Inches(7.1), Inches(2.35), Inches(5.4), Inches(0.4), "NEXT · PHASE IV", 12, CORAL, bold=True)
    add_multiline(s, Inches(7.1), Inches(2.85), Inches(5.4), Inches(3.6), [
        "FastAPI + WebSocket API layer",
        "Node view — SOS button, live delivery status",
        "Rescue dashboard — live topology, SOS queue",
        "Integration across real machines, demo rehearsal",
    ], 14, WHITE, space_after=14)

    # ================================================================ Slide 12: Thank you
    s = blank_slide(prs)
    add_text(s, Inches(0), Inches(3.0), Inches(13.333), Inches(1.0),
              "THANK YOU", 44, WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0), Inches(3.95), Inches(13.333), Inches(0.5),
              "Questions?", 20, CORAL, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0), Inches(5.1), Inches(13.333), Inches(0.4),
              "Miti Shah  ·  Twissha Shah  ·  Tanvi Sheth", 15, CYAN, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0), Inches(5.8), Inches(13.333), Inches(0.35),
              "EcoMesh · Group 20", 11, GREY, align=PP_ALIGN.CENTER)

    prs.save(OUT)
    print("done:", OUT)


if __name__ == "__main__":
    build()
